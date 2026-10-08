import io

import pytest
from docx import Document

from app.docs.extract import UnsupportedFileError, detect_kind, extract_text


def _make_docx(text: str) -> bytes:
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_detect_pdf_by_magic():
    assert detect_kind("algo.pdf", b"%PDF-1.7\n...") == "pdf"


def test_detect_docx_by_magic_and_ext():
    data = _make_docx("hola")
    assert detect_kind("doc.docx", data) == "docx"


def test_reject_unknown():
    with pytest.raises(UnsupportedFileError):
        detect_kind("evil.exe", b"MZ\x90\x00")


def test_extract_docx_text():
    data = _make_docx("Primera línea\nSegunda línea")
    out = extract_text("doc.docx", data)
    assert "Primera línea" in out
    assert "Segunda línea" in out


# --- Tablas y estructuras de .docx ------------------------------------------

from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls  # noqa: E402

from app.docs.extract import extract_docx  # noqa: E402

_W = nsdecls("w")


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _tbl_xml(*cells: str) -> str:
    tcs = "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in cells)
    return f"<w:tbl {_W}><w:tblPr/><w:tblGrid/><w:tr>{tcs}</w:tr></w:tbl>"


def test_docx_tables_in_reading_order():
    doc = Document()
    doc.add_paragraph("Antes")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Misión"
    t.cell(0, 1).text = "Razón de ser"
    doc.add_paragraph("Después")
    assert extract_docx(_save(doc)).split("\n") == ["Antes", "Misión. Razón de ser", "Después"]


def test_docx_cell_separator_respects_punctuation_and_icons():
    doc = Document()
    doc.element.body.insert(0, parse_xml(_tbl_xml("Fase:", "Diagnóstico.", "Fin")))
    doc.element.body.insert(1, parse_xml(_tbl_xml("■", "1. Encuadre")))
    assert extract_docx(_save(doc)).split("\n") == ["Fase: Diagnóstico. Fin", "■ 1. Encuadre"]


def test_docx_nested_table():
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Exterior"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Anidada"
    out = extract_docx(_save(doc))
    assert "Exterior" in out and "Anidada" in out


def test_docx_merged_cell_read_once():
    doc = Document()
    t = doc.add_table(rows=1, cols=3)
    t.cell(0, 0).merge(t.cell(0, 2)).text = "Combinada"
    assert extract_docx(_save(doc)).count("Combinada") == 1


def test_docx_content_control_and_custom_xml():
    doc = Document()
    body = doc.element.body
    body.insert(0, parse_xml(
        f"<w:sdt {_W}><w:sdtPr/><w:sdtContent>{_tbl_xml('EnControl')}"
        "<w:p><w:r><w:t>ParrafoControl</w:t></w:r></w:p></w:sdtContent></w:sdt>"))
    body.insert(1, parse_xml(f'<w:customXml {_W} w:element="x">{_tbl_xml("EnCustom")}</w:customXml>'))
    out = extract_docx(_save(doc))
    assert out.split("\n") == ["EnControl", "ParrafoControl", "EnCustom"]


def test_docx_tracked_changes():
    doc = Document()
    p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]._p
    p.append(parse_xml(f'<w:ins {_W} w:id="1" w:author="x"><w:r><w:t>Insertado</w:t></w:r></w:ins>'))
    p.append(parse_xml(f'<w:del {_W} w:id="2" w:author="x"><w:r><w:delText>Borrado</w:delText></w:r></w:del>'))
    assert extract_docx(_save(doc)) == "Insertado"


def test_docx_text_box_read_once():
    doc = Document()
    p = doc.add_paragraph("Antes")._p
    box = "<w:txbxContent><w:p><w:r><w:t>Cuadro</w:t></w:r></w:p></w:txbxContent>"
    p.append(parse_xml(
        f'<w:r {_W} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent>'
        f'<mc:Choice Requires="wps"><w:drawing><wps:wsp><wps:txbx>{box}</wps:txbx></wps:wsp></w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict><v:shape><v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r>'))
    assert extract_docx(_save(doc)).split("\n") == ["Antes", "Cuadro"]
